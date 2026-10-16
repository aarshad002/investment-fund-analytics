import json
from pathlib import Path
import pytest
import psycopg
from investment_fund_analytics.benchmark_loading import extract
from investment_fund_analytics.database_loading import DEFAULT_MANIFEST,seed

ROOT=Path(__file__).resolve().parents[1]


def test_existing_official_benchmarks_preserve_missing_levels():
    import openpyxl
    manifest=ROOT/DEFAULT_MANIFEST
    if not manifest.exists():pytest.skip('Official raw files are local only')
    for a in json.loads(manifest.read_text()):
        records,missing=extract(ROOT,a)
        workbook=openpyxl.load_workbook(ROOT/a['raw_file'],read_only=True,data_only=True)
        try:
            rows=list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
            header=next(i for i,row in enumerate(rows) if 'Date' in row and 'Index Level' in row)
            date_col=rows[header].index('Date'); value_col=rows[header].index('Index Level')
            expected={}; absent=set()
            from datetime import datetime
            from decimal import Decimal
            for row in rows[header+1:]:
                if not row[date_col]:continue
                day=datetime.strptime(row[date_col],'%d.%m.%Y').date()
                if row[value_col] in (None, 'N/A', ''):absent.add(day)
                else:expected[day]=Decimal(str(row[value_col]))
            assert [(r['observation_date'],r['value']) for r in records]==sorted(expected.items())
            assert set(missing)==absent
            assert len({r['observation_date'] for r in records})==len(records)
            assert all(r['value']>0 for r in records)
        finally:workbook.close()


def test_changed_benchmark_raw_checksum_rejected():
    manifest=ROOT/DEFAULT_MANIFEST
    if not manifest.exists():pytest.skip('Official raw files are local only')
    acquisition=json.loads(manifest.read_text())[0]
    acquisition['file_sha256']='0'*64
    with pytest.raises(ValueError,match='identity'):extract(ROOT,acquisition)


def test_live_assignments_reject_overlapping_periods(isolated_database):
    conn = isolated_database
    # Real constraint DDL; production application grants are irrelevant to this test.
    ddl = (ROOT/'sql/benchmark_schema.sql').read_text(encoding='utf-8')
    conn.execute('\n'.join(line for line in ddl.splitlines() if not line.startswith('GRANT ')))
    fund,_ = seed(conn, {'isin':'LU0380865021'})
    index = conn.execute('''INSERT INTO core.instrument(instrument_code,instrument_name,instrument_type,
                      value_basis,currency,adjustment_status)
                      VALUES('TEST:INDEX','Synthetic constraint-test index','index','total_return_index','EUR','total_return')
                      RETURNING instrument_id''').fetchone()[0]
    benchmark = conn.execute('''INSERT INTO core.benchmark(benchmark_code,instrument_id,benchmark_name,
                      identifier_notes,benchmark_kind,return_basis,metadata_url,date_alignment_status,limitations)
                      VALUES('TEST:BENCHMARK',%s,'Synthetic constraint-test benchmark','Internal test identity',
                      'official_index','net_total_return','https://example.invalid','unverified','Test-only metadata')
                      RETURNING benchmark_id''',(index,)).fetchone()[0]
    conn.execute('''INSERT INTO core.benchmark_assignment(fund_instrument_id,benchmark_id,
                   effective_start_date,effective_end_date,rationale)
                   VALUES(%s,%s,'2021-09-30','2026-09-30','Test-only assignment')''',(fund,benchmark))
    before = conn.execute('SELECT * FROM core.benchmark_assignment').fetchall()
    with pytest.raises(psycopg.errors.ExclusionViolation),conn.transaction():
        conn.execute('''INSERT INTO core.benchmark_assignment(fund_instrument_id,benchmark_id,effective_start_date,effective_end_date,rationale)
                     VALUES(%s,%s,'2022-01-01','2022-12-31','isolated rollback constraint check')''',(fund,benchmark))
    assert conn.execute('SELECT * FROM core.benchmark_assignment').fetchall() == before


def test_live_benchmark_guard_and_interval_metadata(readonly_project_database):
    conn = readonly_project_database
    with conn.transaction():
        if not conn.execute("SELECT to_regclass('analytics.nav_intervals')").fetchone()[0]:
            pytest.skip('Controlled interval view not applied')
        assert conn.execute("SELECT count(*) FROM core.benchmark WHERE date_alignment_status='unverified'").fetchone()[0]==2
        assert conn.execute("SELECT count(*) FROM core.observation o JOIN core.benchmark b USING(instrument_id) WHERE quality_status<>'flagged'").fetchone()[0]==0
        assert conn.execute("SELECT count(*) FROM core.observation o JOIN core.benchmark b USING(instrument_id)").fetchone()[0]>0
        assert conn.execute("SELECT count(*) FROM analytics.nav_intervals WHERE policy_status<>'approved_controlled_nav_v1'").fetchone()[0]==0
        assert conn.execute("SELECT min(interval_days) FROM analytics.nav_intervals WHERE interval_days IS NOT NULL").fetchone()[0]>=1


def test_live_sql_index_rows_match_python_raw_extraction(readonly_project_database):
    manifest=ROOT/DEFAULT_MANIFEST
    if not manifest.exists() or not (ROOT/'.env.app').exists():pytest.skip('Local raw data and app connection needed')
    from investment_fund_analytics.benchmark_loading import BENCHMARKS
    conn = readonly_project_database
    with conn.transaction():
        for a in json.loads(manifest.read_text()):
            records,_=extract(ROOT,a)
            stored=conn.execute('''SELECT observation_date,value,raw_row_number FROM core.observation o
                 JOIN core.instrument i USING(instrument_id) WHERE i.instrument_code=%s ORDER BY observation_date''',
                 (BENCHMARKS[a['isin']]['code'],)).fetchall()
            assert stored==[(r['observation_date'],r['value'],r['raw_row_number']) for r in records]
