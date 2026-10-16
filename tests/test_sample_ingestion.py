from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from investment_fund_analytics.sample_ingestion import (
    INSTRUMENTS, WorkbookStructureError, parse_workbook,
)

META = INSTRUMENTS['LU0380865021']


def make_workbook(tmp_path, rows=None, currency='EUR', header_row=14, isin=META.isin):
    book = Workbook()
    sheet = book.active
    sheet.title = isin
    for number, label, value in [(2, 'Name', META.fund_name), (3, 'ISIN', isin),
                                  (4, 'Currency', currency), (5, 'As of', '07.10.2026')]:
        sheet.cell(number, 1, label)
        sheet.cell(number, 2, value)
    for column, value in enumerate(['Date', 'Asset Under Management', 'NAV', 'Index Level', 'Interim Profits'], 1):
        sheet.cell(header_row, column, value)
    for number, (day, nav) in enumerate(rows if rows is not None else [('02.10.2023', 100), ('03.10.2023', 101)], header_row+1):
        for column, value in [(1, day), (2, 123456), (3, nav), (4, 'N/A'), (5, 'N/A')]:
            sheet.cell(number, column, value)
    path = tmp_path / 'sample.xlsx'
    book.save(path)
    return path


def codes(result):
    return set(result.issues.check_code)


def test_normal_parsing_sorted_and_metadata(tmp_path):
    result = parse_workbook(make_workbook(tmp_path, [('03.10.2023', 101), ('02.10.2023', '100.25')]), META)
    assert len(result.clean) == 2
    assert result.clean.observation_date.is_monotonic_increasing
    assert result.clean.value.tolist() == [Decimal('100.25'), Decimal('101')]
    assert result.clean.raw_row_number.tolist() == [16, 15]
    assert result.clean.instrument_code.unique().tolist() == [META.instrument_code]
    assert result.clean.source_code.unique().tolist() == ['S01']
    assert result.structure['input_order'] == 'descending'
    assert result.issues.empty


def test_header_discovery_and_excel_date(tmp_path):
    result = parse_workbook(make_workbook(tmp_path, [(date(2023, 10, 2), 100)], header_row=8), META)
    assert result.structure['header_row'] == 8
    assert result.clean.observation_date.iloc[0].date() == date(2023, 10, 2)


@pytest.mark.parametrize('second', [100, 101])
def test_duplicate_dates_preserved_and_quarantined(tmp_path, second):
    result = parse_workbook(make_workbook(tmp_path, [('02.10.2023', 100), ('02.10.2023', second)]), META)
    assert len(result.records) == 2
    assert result.clean.empty
    assert result.issues.check_code.tolist() == ['duplicate_date', 'duplicate_date']


@pytest.mark.parametrize('value', ['garbage', 'NaN', 'Infinity', True, '=1+1'])
def test_invalid_numeric(tmp_path, value):
    result = parse_workbook(make_workbook(tmp_path, [('02.10.2023', value)]), META)
    assert 'invalid_numeric' in codes(result)
    assert result.clean.empty
    assert len(result.records) == 1


@pytest.mark.parametrize('value', [0, -1])
def test_nonpositive_nav(tmp_path, value):
    result = parse_workbook(make_workbook(tmp_path, [('02.10.2023', value)]), META)
    assert 'nonpositive_nav' in codes(result)
    assert result.clean.empty


@pytest.mark.parametrize('value', [None, 'N/A'])
def test_missing_nav(tmp_path, value):
    result = parse_workbook(make_workbook(tmp_path, [('02.10.2023', value)]), META)
    assert 'missing_nav' in codes(result)
    assert result.clean.empty


@pytest.mark.parametrize('day,code', [(None, 'missing_date'), ('31.02.2023', 'invalid_date'), ('2023-10-02', 'invalid_date')])
def test_missing_or_invalid_date(tmp_path, day, code):
    result = parse_workbook(make_workbook(tmp_path, [(day, 100)]), META)
    assert code in codes(result)
    assert result.clean.empty


def test_unexpected_currency(tmp_path):
    result = parse_workbook(make_workbook(tmp_path, currency='USD'), META)
    assert 'unexpected_currency' in codes(result)
    assert result.clean.empty
    assert result.records.currency.tolist() == ['USD', 'USD']


def test_unexpected_identity(tmp_path):
    result = parse_workbook(make_workbook(tmp_path, isin='LU0290355717'), META)
    assert 'identity_mismatch' in codes(result)
    assert result.clean.empty


def test_extreme_change_retained_flagged(tmp_path):
    result = parse_workbook(make_workbook(tmp_path, [('02.10.2023', 100), ('03.10.2023', 150)]), META)
    assert len(result.clean) == 2
    assert result.clean.quality_status.tolist() == ['valid', 'flagged']
    assert 'extreme_change' in codes(result)


def test_invalid_row_does_not_create_jump_comparison(tmp_path):
    result = parse_workbook(make_workbook(tmp_path, [('02.10.2023', 100), ('03.10.2023', None), ('04.10.2023', 200)]), META)
    assert 'extreme_change' not in codes(result)
    assert len(result.clean) == 2


def test_unordered_dates_flagged_and_sorted(tmp_path):
    result = parse_workbook(make_workbook(tmp_path, [('03.10.2023', 101), ('02.10.2023', 100), ('04.10.2023', 102)]), META)
    assert 'unordered_dates' in codes(result)
    assert result.clean.observation_date.is_monotonic_increasing


def test_missing_header_fails_explicitly(tmp_path):
    book = Workbook()
    path = tmp_path / 'no_header.xlsx'
    book.save(path)
    with pytest.raises(WorkbookStructureError, match='found 0'):
        parse_workbook(path, META)


@pytest.mark.parametrize('isin', sorted(INSTRUMENTS))
@pytest.mark.parametrize('year', [2021, 2023, 2026])
def test_real_local_sample(isin, year):
    path = Path(__file__).resolve().parents[1] / 'data/raw/phase0' / f'{isin}_{year}_sample.xlsx'
    if not path.exists():
        pytest.skip('Raw local samples are intentionally excluded from Git')
    result = parse_workbook(path, INSTRUMENTS[isin])
    assert len(result.records) == len(result.clean) == 5
    assert result.structure['sheet_names'] == [isin]
    assert result.structure['header_row'] == 14
    assert result.issues.empty
    assert result.clean.observation_date.is_monotonic_increasing
    assert result.structure['optional_missing']['Index Level'] == (3 if year == 2026 else 0)
